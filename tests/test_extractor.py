"""Tests for the DOCX extractor (T1.1).

Covers: body order (tables interleaved), block classification, text cleaning
(NBSP, field codes, hyperlinks), table cell meta (incl. empty cells and
gridSpan), sequential ids/order/word count, deterministic doc_id, and the
typed errors for encrypted / corrupt / missing files.
"""
from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from lexaudit.extractor.parse import (
    CorruptDocxError,
    EncryptedDocxError,
    parse_docx,
)

FIXTURES = Path(__file__).resolve().parent / "fixtures" / "docx"
# Synthetic T0.2 fixtures only (docx/ may also hold real contracts).
SYNTHETIC = ["tourism", "labor", "lease", "clean"]


# --- helpers ------------------------------------------------------------------


def _texts(doc) -> list[str]:
    return [b.text for b in doc.blocks]


def _index_of(texts: list[str], substring: str) -> int:
    for i, t in enumerate(texts):
        if substring in t:
            return i
    raise AssertionError(f"not found: {substring!r}")


@pytest.fixture(scope="module")
def fixtures():
    """Parse the four synthetic T0.2 fixtures once per module run."""
    return {name: parse_docx(str(FIXTURES / f"{name}.docx")) for name in SYNTHETIC}


# --- T0.2 fixtures -------------------------------------------------------------


def test_fixtures_present(fixtures):
    assert set(fixtures) == set(SYNTHETIC)


def test_doc_id_deterministic_and_unique(fixtures):
    ids = [d.doc_id for d in fixtures.values()]
    assert len(set(ids)) == len(SYNTHETIC)
    for doc in fixtures.values():
        assert doc.doc_id.startswith("d_")
        # re-parsing the same file yields the same id
        assert parse_docx(str(FIXTURES / f"{doc.filename}")).doc_id == doc.doc_id


def test_block_ids_sequential_and_order_matches(fixtures):
    for name, doc in fixtures.items():
        for i, block in enumerate(doc.blocks):
            assert block.block_id == f"b_{i + 1:04d}", name
            assert block.meta.order == i, name
        assert doc.stats.total_blocks == len(doc.blocks), name


def test_stats_words(fixtures):
    for doc in fixtures.values():
        expected = sum(len(b.text.split()) for b in doc.blocks)
        assert doc.stats.words == expected


def test_tourism_table_blocks(fixtures):
    doc = fixtures["tourism"]
    cells = [b for b in doc.blocks if b.type == "table_cell"]
    # 5 program rows + header row + total row = 7 rows x 3 cols, no empty cells
    assert len(cells) == 21
    tables = {b.meta.table for b in cells}
    assert tables == {"table_1"}
    # header row values and a data cell
    assert "Стоимость, тенге" in _texts(doc)
    # cells carry row/col, paragraphs don't
    for b in doc.blocks:
        if b.type == "table_cell":
            assert b.meta.row is not None and b.meta.col is not None
        else:
            assert b.meta.table is None and b.meta.row is None and b.meta.col is None


def test_tourism_headings_and_lists(fixtures):
    doc = fixtures["tourism"]
    headings = [b for b in doc.blocks if b.type == "heading"]
    # 6 numbered section headings; the doc title uses the "Title" style,
    # which is not a "Heading"/"Заголовок" style, so it is a paragraph
    assert len(headings) == 6
    assert all(b.meta.style and "Heading" in b.meta.style for b in headings)
    list_items = [b for b in doc.blocks if b.type == "list_item"]
    # 5 numbered program-day items
    assert len(list_items) == 5


def test_tourism_table_interleaved_in_body_order(fixtures):
    """The table must sit between its surrounding paragraphs (T1.1 rule 1)."""
    doc = fixtures["tourism"]
    texts = _texts(doc)
    i_last_program_item = _index_of(texts, "День 5 (09.11.2026)")
    i_before = _index_of(texts, "Состав мероприятий и цены по программе тура:")
    i_total = _index_of(texts, "Итого")
    i_next = _index_of(texts, "3.1. Туроператор обязан оказать услуги")
    # program list items, then caption, then the table itself, then the next section
    assert i_last_program_item < i_before < i_total < i_next
    # every table block must carry the table meta
    total_cell = next(b for b in doc.blocks if b.text == "Итого")
    assert total_cell.type == "table_cell"
    assert total_cell.meta.table is not None


# --- synthetic .docx documents (t11_docs) --------------------------------------


def _make(doc_path: Path, builder) -> Path:
    doc = Document()
    builder(doc)
    doc_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(doc_path)
    return doc_path


def test_empty_paragraphs_dropped_empty_cells_kept(tmp_path):
    def build(doc: Document) -> None:
        doc.add_paragraph("Первый абзац.")
        doc.add_paragraph()  # empty spacer paragraph
        t = doc.add_table(rows=2, cols=2)
        t.style = "Table Grid"
        t.rows[0].cells[0].text = "Заполнен"
        # t.rows[0].cells[1] and row 1 left empty
        doc.add_paragraph("Последний абзац.")

    result = parse_docx(str(_make(tmp_path / "gap.docx", build)))
    texts = _texts(result)
    assert "" not in [b.text for b in result.blocks if b.type != "table_cell"]
    cells = [b for b in result.blocks if b.type == "table_cell"]
    assert len(cells) == 4
    empty = [c for c in cells if c.text == ""]
    assert len(empty) == 3
    assert texts[0] == "Первый абзац."
    assert texts[-1] == "Последний абзац."
    assert result.stats.words == sum(len(t.split()) for t in texts)


def test_nested_table_flattened(tmp_path):
    def build(doc: Document) -> None:
        outer = doc.add_table(rows=1, cols=1)
        outer.style = "Table Grid"
        outer.rows[0].cells[0].text = "Внешняя"
        inner = outer.rows[0].cells[0].add_table(rows=1, cols=1)
        inner.style = "Table Grid"
        inner.rows[0].cells[0].text = "Вложенная"

    result = parse_docx(str(_make(tmp_path / "nested.docx", build)))
    cells = [b for b in result.blocks if b.type == "table_cell"]
    assert {c.text for c in cells} == {"Внешняя", "Вложенная"}
    assert {c.meta.table for c in cells} == {"table_1", "table_2"}


def test_grid_span_merging(tmp_path):
    def build(doc: Document) -> None:
        t = doc.add_table(rows=1, cols=2)
        t.style = "Table Grid"
        # merge the two cells of the single row
        t.rows[0].cells[0].merge(t.rows[0].cells[1]).text = "Слитная ячейка"

    result = parse_docx(str(_make(tmp_path / "span.docx", build)))
    cells = [b for b in result.blocks if b.type == "table_cell"]
    assert len(cells) == 1
    assert cells[0].text == "Слитная ячейка"
    assert cells[0].meta.row == 0
    assert cells[0].meta.col == 0


def test_cleaning_nbsp_and_whitespace(tmp_path):
    def build(doc: Document) -> None:
        p = doc.add_paragraph()
        p.add_run("Договор\u00a0№\u00a014")  # nbsp around №
        p.add_run("\n  \t  \u202fпробелы")

    result = parse_docx(str(_make(tmp_path / "nbsp.docx", build)))
    assert _texts(result) == ["Договор № 14 пробелы"]


def test_field_codes_and_hyperlinks(tmp_path):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement

    def add_hyperlink(paragraph, url: str, text: str) -> None:
        r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), r_id)
        run = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = text
        run.append(t)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)

    def build(doc: Document) -> None:
        p = doc.add_paragraph()
        p.add_run("Ссылка: ")
        add_hyperlink(p, "https://adilet.zan.kz", "закон РК")  # visible text kept
        p.add_run(" и дальше текст.")
        # DATE field: instruction (must be dropped) + cached result (kept)
        field = doc.add_paragraph()
        r = field.add_run()
        f_begin = r._r.makeelement(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldChar",
            {"{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldCharType": "begin"},
        )
        instr = r._r.makeelement(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}instrText", {}
        )
        instr.text = " DATE \\@ \"dd.MM.yyyy\" "
        f_sep = r._r.makeelement(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldChar",
            {"{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldCharType": "separate"},
        )
        t = r._r.makeelement(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t", {}
        )
        t.text = "01.01.2026"
        f_end = r._r.makeelement(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldChar",
            {"{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldCharType": "end"},
        )
        for el in (f_begin, instr, f_sep, t, f_end):
            r._r.append(el)

    result = parse_docx(str(_make(tmp_path / "fields.docx", build)))
    texts = _texts(result)
    assert texts[0] == "Ссылка: закон РК и дальше текст."
    assert texts[1] == "01.01.2026"


def test_style_based_lists_and_russian_heading(tmp_path):
    def build(doc: Document) -> None:
        doc.add_heading("Раздел 1", level=1)
        doc.add_paragraph("пункт один", style="List Bullet")
        doc.add_paragraph("1. пунктированный текст")
        doc.add_paragraph("а) пункт с литерой")
        doc.add_paragraph("• буллет")
        doc.add_paragraph("обычный текст без маркера")

    result = parse_docx(str(_make(tmp_path / "styles.docx", build)))
    types = [b.type for b in result.blocks]
    assert types == [
        "heading",
        "list_item",
        "list_item",
        "list_item",
        "list_item",
        "paragraph",
    ]


def test_non_utf8_name_survives_json_roundtrip(tmp_path):
    def build(doc: Document) -> None:
        doc.add_paragraph("Договор на оказание туристических услуг")

    path = _make(tmp_path / "турсервис (копия).docx", build)
    result = parse_docx(str(path))
    import json

    payload = json.dumps(result.model_dump(), ensure_ascii=False)
    assert "турсервис (копия).docx" in payload
    assert result.filename == "турсервис (копия).docx"


# --- typed errors ---------------------------------------------------------------


def test_missing_file_raises_corrupt(tmp_path):
    with pytest.raises(CorruptDocxError, match="file not found"):
        parse_docx(str(tmp_path / "nope.docx"))


def test_not_a_zip_raises_corrupt(tmp_path):
    path = tmp_path / "fake.docx"
    path.write_bytes(b"this is definitely not a docx file")
    with pytest.raises(CorruptDocxError, match="corrupt"):
        parse_docx(str(path))


def test_zip_without_ooffice_parts_raises_corrupt(tmp_path):
    import zipfile

    path = tmp_path / "plain.docx"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("hello.txt", "not office open xml")
    with pytest.raises(CorruptDocxError, match="corrupt"):
        parse_docx(str(path))


def test_ole2_magic_raises_encrypted(tmp_path):
    path = tmp_path / "locked.docx"
    # OLE2 compound file header — what an encrypted OOXML document starts with.
    path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 32)
    with pytest.raises(EncryptedDocxError, match="password-encrypted"):
        parse_docx(str(path))
